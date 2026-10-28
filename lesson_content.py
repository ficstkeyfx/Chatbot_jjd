from __future__ import annotations

import base64
import io
import os
import re
import shutil
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any


MAX_FILE_BYTES = int(os.getenv("SUMMARY_MAX_FILE_BYTES", str(10 * 1024 * 1024)))
MAX_TEXT_CHARS = int(os.getenv("SUMMARY_MAX_CHARS", "50000"))
MAX_PDF_PAGES = int(os.getenv("SUMMARY_MAX_PDF_PAGES", "20"))
MAX_IMAGE_PIXELS = int(os.getenv("SUMMARY_MAX_IMAGE_PIXELS", "20000000"))
MAX_DOCX_UNCOMPRESSED_BYTES = int(
    os.getenv("SUMMARY_MAX_DOCX_UNCOMPRESSED_BYTES", str(50 * 1024 * 1024))
)
OCR_LANG = os.getenv("OCR_LANG", "vie+eng").strip() or "vie+eng"
OCR_MIN_PAGE_CHARS = int(os.getenv("OCR_MIN_PAGE_CHARS", "30"))

MEDIA_TYPE_KEYS = (
    "file_type",
    "fileType",
    "format",
    "type",
    "media_type",
    "attachment_type",
)
MEDIA_URL_KEYS = (
    "document_url",
    "documentUrl",
    "file_url",
    "fileUrl",
    "download_url",
    "downloadUrl",
    "photo_url",
    "photoUrl",
    "photo",
    "image_url",
    "imageUrl",
    "media_url",
    "mediaUrl",
    "original_url",
    "originalUrl",
    "source_url",
    "sourceUrl",
    "href",
)


class LessonContentError(RuntimeError):
    """A safe, user-facing content extraction error."""


@dataclass(frozen=True)
class IncomingLesson:
    kind: str
    url: str
    filename: str
    mime_type: str
    size: int | None = None
    # Dữ liệu tải lên trực tiếp từ giao diện web; khi có thì không cần tải từ URL.
    data: bytes | None = None


def _message_payload(update: Any) -> dict[str, Any]:
    message = getattr(update, "message", None)
    update_payload = dict(getattr(update, "api_kwargs", {}) or {})
    message_payload = dict(getattr(message, "api_kwargs", {}) or {})
    serialized_payload: dict[str, Any] = {}
    to_dict = getattr(update, "to_dict", None)
    if callable(to_dict):
        try:
            serialized = to_dict()
            if isinstance(serialized, dict):
                serialized_payload = serialized
        except (AttributeError, TypeError, ValueError):
            pass
    return {
        "update": update_payload,
        "message": message_payload,
        "serialized": serialized_payload,
    }


def extract_image_caption(update: Any) -> str:
    """Read the documented Zalo image caption, which this SDK keeps in api_kwargs."""
    message = getattr(update, "message", None)
    if message is None:
        return ""
    raw_message = getattr(message, "api_kwargs", {}) or {}
    raw_update = getattr(update, "api_kwargs", {}) or {}
    update_message = raw_update.get("message") if isinstance(raw_update, Mapping) else None
    candidates = (
        getattr(message, "caption", None),
        raw_message.get("caption") if isinstance(raw_message, Mapping) else None,
        update_message.get("caption") if isinstance(update_message, Mapping) else None,
        getattr(message, "text", None),
    )
    for candidate in candidates:
        if isinstance(candidate, str) and candidate.strip():
            return candidate.strip()
    return ""


def _walk_dicts(value: Any):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _walk_dicts(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_dicts(child)


def _walk_media_dicts(value: Any, inherited_type: str = ""):
    """Yield payload mappings together with their nearest attachment type."""
    if isinstance(value, dict):
        declared_type = str(
            _first(
                value,
                MEDIA_TYPE_KEYS,
                inherited_type,
            )
        )
        yield value, declared_type
        for child in value.values():
            yield from _walk_media_dicts(child, declared_type)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_media_dicts(child, inherited_type)


def _first(mapping: dict[str, Any], names: tuple[str, ...], default=None):
    for name in names:
        value = mapping.get(name)
        if value not in (None, ""):
            return value
    return default


def _url_value(value: Any) -> str | None:
    if isinstance(value, str):
        candidate = value.strip()
        return candidate or None
    if isinstance(value, dict):
        direct = _first(value, (*MEDIA_URL_KEYS, "url"))
        if direct is not None and direct is not value:
            candidate = _url_value(direct)
            if candidate:
                return candidate
        for child in value.values():
            candidate = _url_value(child)
            if candidate and candidate.startswith(("https://", "http://")):
                return candidate
    if isinstance(value, (list, tuple)):
        for child in reversed(value):
            candidate = _url_value(child)
            if candidate:
                return candidate
    return None


def _kind_from_metadata(
    filename: str,
    mime_type: str,
    declared_type: str = "",
) -> str | None:
    extension = Path(urllib.parse.urlparse(filename).path).suffix.lower()
    mime = mime_type.lower().split(";", 1)[0].strip()
    declared = declared_type.lower().lstrip(".").strip()
    if mime == "application/pdf" or extension == ".pdf" or declared == "pdf":
        return "pdf"
    if (
        mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or extension == ".docx"
        or declared == "docx"
    ):
        return "docx"
    if mime in {
        "image/jpeg",
        "image/png",
        "image/webp",
        "image/heic",
        "image/heif",
        "image/tiff",
        "image/bmp",
        "image/gif",
    } or extension in {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
        ".heic",
        ".heif",
        ".tif",
        ".tiff",
        ".bmp",
        ".gif",
    } or declared in {
        "jpg",
        "jpeg",
        "png",
        "webp",
        "heic",
        "heif",
        "tif",
        "tiff",
        "bmp",
        "gif",
        "image",
        "photo",
    }:
        return "image"
    return None


def _has_attachment_metadata(filename: str, mime_type: str, media_type: str) -> bool:
    return bool(
        filename
        or mime_type
        or media_type.lower()
        in {"document", "file", "attachment", "photo", "image"}
    )


def extract_lesson_source(update: Any) -> IncomingLesson | None:
    message = getattr(update, "message", None)
    if not message:
        return None

    photo_url = _url_value(getattr(message, "photo_url", None))
    if photo_url and photo_url.startswith("https://"):
        return IncomingLesson("image", photo_url, "lesson.jpg", "image/jpeg")

    payload = _message_payload(update)
    for item, inherited_type in _walk_media_dicts(payload):
        mime_type = str(
            _first(item, ("mime_type", "mimetype", "content_type", "contentType"), "")
        )
        filename = str(
            _first(
                item,
                ("file_name", "fileName", "filename", "name", "original_name"),
                "",
            )
        )
        media_type = str(
            _first(
                item,
                MEDIA_TYPE_KEYS,
                inherited_type,
            )
        )
        if mime_type.lower().startswith("audio/") or media_type.lower() in {"audio", "voice", "voice_note"}:
            continue
        url = _url_value(_first(item, MEDIA_URL_KEYS))
        if not media_type and any(key in item for key in ("photo", "photo_url", "photoUrl")):
            media_type = "image"
        has_attachment_metadata = _has_attachment_metadata(
            filename, mime_type, media_type
        )
        if url is None and has_attachment_metadata:
            url = _url_value(item.get("url"))
        if url is None:
            continue
        if not url.startswith("https://"):
            continue
        inferred_name = filename or urllib.parse.unquote(Path(urllib.parse.urlparse(url).path).name)
        kind = _kind_from_metadata(inferred_name, mime_type, media_type)
        if not kind and not has_attachment_metadata:
            continue
        size = _first(item, ("file_size", "size", "bytes"))
        parsed_size = int(str(size)) if str(size).isdigit() else None
        return IncomingLesson(
            kind=kind or "auto",
            url=url,
            filename=inferred_name or f"lesson.{kind or 'bin'}",
            mime_type=mime_type,
            size=parsed_size,
        )
    return None


def safe_lesson_metadata(update: Any) -> dict[str, Any]:
    """Return only structural metadata for diagnosing unsupported Zalo payloads."""
    payload = _message_payload(update)
    nested_keys: set[str] = set()
    declared_types: set[str] = set()
    event_names: set[str] = set()
    for item in _walk_dicts(payload):
        nested_keys.update(str(key) for key in item.keys())
        event_name = item.get("event_name")
        if isinstance(event_name, str) and event_name:
            event_names.add(event_name[:100])
        for key in ("type", "file_type", "fileType", "format", "media_type"):
            value = item.get(key)
            if isinstance(value, str) and value:
                declared_types.add(value[:50])
    return {
        "message_type": str(getattr(getattr(update, "message", None), "message_type", "")),
        "event_names": sorted(event_names),
        "keys": sorted(nested_keys),
        "declared_types": sorted(declared_types),
    }


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def validate_text_length(text: str) -> str:
    normalized = normalize_text(text)
    if not normalized:
        raise LessonContentError("Không tìm thấy nội dung chữ để tóm tắt.")
    if len(normalized) > MAX_TEXT_CHARS:
        raise LessonContentError(
            f"Nội dung dài quá {MAX_TEXT_CHARS:,} ký tự. Hãy chia bài học thành nhiều phần nhỏ."
        )
    return normalized


def download_lesson(source: IncomingLesson) -> bytes:
    if source.data is not None:
        if len(source.data) > MAX_FILE_BYTES:
            raise LessonContentError("Tệp vượt quá giới hạn 10 MB.")
        return source.data
    if source.size is not None and source.size > MAX_FILE_BYTES:
        raise LessonContentError("Tệp vượt quá giới hạn 10 MB.")
    parsed = urllib.parse.urlparse(source.url)
    if parsed.scheme != "https" or not parsed.hostname:
        raise LessonContentError("URL tệp phải dùng HTTPS.")
    request = urllib.request.Request(
        source.url,
        headers={"User-Agent": "zalo-health-bot/1.0"},
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            declared = response.headers.get("Content-Length")
            if declared and int(declared) > MAX_FILE_BYTES:
                raise LessonContentError("Tệp vượt quá giới hạn 10 MB.")
            data = response.read(MAX_FILE_BYTES + 1)
    except LessonContentError:
        raise
    except (urllib.error.URLError, ValueError) as exc:
        reason = getattr(exc, "reason", str(exc))
        raise LessonContentError(f"Không tải được tệp bài học: {reason}") from exc
    if len(data) > MAX_FILE_BYTES:
        raise LessonContentError("Tệp vượt quá giới hạn 10 MB.")
    return data


def _require_tesseract():
    try:
        import pytesseract
    except ImportError as exc:
        raise LessonContentError("Thiếu thư viện pytesseract. Hãy cài lại requirements.txt.") from exc

    configured = os.getenv("TESSERACT_CMD", "").strip()
    executable = configured or shutil.which("tesseract")
    if not executable or not Path(executable).is_file():
        raise LessonContentError(
            "Chưa tìm thấy Tesseract OCR. Hãy cài Tesseract và cấu hình TESSERACT_CMD."
        )
    pytesseract.pytesseract.tesseract_cmd = executable
    return pytesseract


def _poppler_path() -> str | None:
    configured = os.getenv("POPPLER_BIN", "").strip()
    if configured:
        executable = Path(configured) / "pdftoppm.exe"
        if not executable.is_file():
            raise LessonContentError(
                "POPPLER_BIN không chứa pdftoppm.exe. Hãy kiểm tra lại cấu hình Poppler."
            )
        return configured
    if shutil.which("pdftoppm"):
        return None
    raise LessonContentError(
        "Chưa tìm thấy Poppler. Hãy cài Poppler và cấu hình POPPLER_BIN để đọc PDF scan."
    )


def _ocr_image(image) -> str:
    pytesseract = _require_tesseract()
    try:
        return pytesseract.image_to_string(image, lang=OCR_LANG)
    except Exception as exc:
        # pytesseract có nhiều lớp lỗi phụ thuộc phiên bản. Chuyển toàn bộ lỗi
        # từ tiến trình OCR thành thông báo an toàn, không làm lộ command/path.
        if exc.__class__.__module__.startswith("pytesseract"):
            raise LessonContentError(
                f"OCR thất bại. Hãy kiểm tra Tesseract và gói ngôn ngữ {OCR_LANG}."
            ) from exc
        raise


def _prepare_ocr_image(image):
    """Normalize contrast and enlarge small screenshots before local OCR."""
    from PIL import Image, ImageOps

    prepared = ImageOps.grayscale(image)
    prepared = ImageOps.autocontrast(prepared)
    longest_edge = max(prepared.size)
    if 0 < longest_edge < 1600:
        scale = min(2.0, 1600 / longest_edge)
        prepared = prepared.resize(
            (
                max(1, round(prepared.width * scale)),
                max(1, round(prepared.height * scale)),
            ),
            Image.Resampling.LANCZOS,
        )
    return prepared


def _extract_image(data: bytes) -> str:
    try:
        from PIL import Image, ImageOps, UnidentifiedImageError
    except ImportError as exc:
        raise LessonContentError("Thiếu Pillow. Hãy cài lại requirements.txt.") from exc
    try:
        from pillow_heif import register_heif_opener
    except ImportError as exc:
        raise LessonContentError("Thiếu pillow-heif. Hãy cài lại requirements.txt.") from exc
    register_heif_opener()
    try:
        with Image.open(io.BytesIO(data)) as opened:
            image = ImageOps.exif_transpose(opened)
            image.load()
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise LessonContentError("Ảnh vượt quá giới hạn 20 megapixel.")
            text = _ocr_image(_prepare_ocr_image(image))
    except LessonContentError:
        raise
    except (UnidentifiedImageError, OSError) as exc:
        raise LessonContentError("Tệp ảnh không hợp lệ hoặc đã bị hỏng.") from exc
    return validate_text_length(text)


def image_data_url(source: IncomingLesson, maximum_edge: int = 2048) -> str:
    """Download and normalize an incoming image for a vision-capable model."""
    try:
        from PIL import Image, ImageOps, UnidentifiedImageError
    except ImportError as exc:
        raise LessonContentError("Thiếu Pillow. Hãy cài lại requirements.txt.") from exc
    try:
        from pillow_heif import register_heif_opener
    except ImportError as exc:
        raise LessonContentError("Thiếu pillow-heif. Hãy cài lại requirements.txt.") from exc

    register_heif_opener()
    data = download_lesson(source)
    try:
        with Image.open(io.BytesIO(data)) as opened:
            image = ImageOps.exif_transpose(opened)
            image.load()
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise LessonContentError("Ảnh vượt quá giới hạn 20 megapixel.")
            image = image.convert("RGB")
            image.thumbnail((maximum_edge, maximum_edge), Image.Resampling.LANCZOS)
            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=88, optimize=True)
    except LessonContentError:
        raise
    except (UnidentifiedImageError, OSError) as exc:
        raise LessonContentError("Tệp ảnh không hợp lệ hoặc đã bị hỏng.") from exc

    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/jpeg;base64,{encoded}"


def _is_docx(data: bytes) -> bool:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return "word/document.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False


def _is_image(data: bytes) -> bool:
    try:
        from PIL import Image, UnidentifiedImageError
    except ImportError as exc:
        raise LessonContentError("Thiếu Pillow. Hãy cài lại requirements.txt.") from exc
    try:
        from pillow_heif import register_heif_opener
    except ImportError as exc:
        raise LessonContentError("Thiếu pillow-heif. Hãy cài lại requirements.txt.") from exc
    register_heif_opener()
    try:
        with Image.open(io.BytesIO(data)) as image:
            return image.format is not None
    except (UnidentifiedImageError, OSError):
        return False


def _extract_text_file(data: bytes) -> str:
    encodings = ("utf-8-sig",)
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ("utf-16", *encodings)
    for encoding in encodings:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        control_count = sum(
            1 for char in text if ord(char) < 32 and char not in "\n\r\t"
        )
        if control_count > max(2, len(text) // 100):
            continue
        return validate_text_length(text)
    raise LessonContentError(
        "Tệp không phải PDF, DOCX, ảnh hoặc văn bản UTF-8/UTF-16 có thể đọc."
    )


def _validate_docx_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                raise LessonContentError("Tệp không phải DOCX hợp lệ.")
            total_size = sum(info.file_size for info in archive.infolist())
            if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise LessonContentError("DOCX giải nén vượt quá giới hạn an toàn 50 MB.")
    except zipfile.BadZipFile as exc:
        raise LessonContentError("Tệp không phải DOCX hợp lệ.") from exc


def _extract_docx(data: bytes) -> str:
    _validate_docx_archive(data)
    try:
        from docx import Document
    except ImportError as exc:
        raise LessonContentError("Thiếu python-docx. Hãy cài lại requirements.txt.") from exc
    try:
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    parts.append(" | ".join(values))
    except (ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise LessonContentError("Không thể đọc nội dung DOCX.") from exc
    return validate_text_length("\n".join(parts))


def _extract_pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise LessonContentError("Tệp không phải PDF hợp lệ.")
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:
        raise LessonContentError("Thiếu pypdf. Hãy cài lại requirements.txt.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise LessonContentError("PDF được bảo vệ bằng mật khẩu nên không thể đọc.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise LessonContentError(f"PDF vượt quá giới hạn {MAX_PDF_PAGES} trang.")

        page_texts: list[str] = []
        pages_needing_ocr: list[int] = []
        for index, page in enumerate(reader.pages, start=1):
            extracted = normalize_text(page.extract_text() or "")
            page_texts.append(extracted)
            try:
                has_images = len(page.images) > 0
            except Exception:
                has_images = False
            if len(extracted) < OCR_MIN_PAGE_CHARS and has_images:
                pages_needing_ocr.append(index)
    except LessonContentError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise LessonContentError("PDF không hợp lệ hoặc đã bị hỏng.") from exc

    if pages_needing_ocr:
        try:
            from pdf2image import convert_from_bytes
        except ImportError as exc:
            raise LessonContentError("Thiếu pdf2image. Hãy cài lại requirements.txt.") from exc
        poppler_path = _poppler_path()
        for page_number in pages_needing_ocr:
            try:
                images = convert_from_bytes(
                    data,
                    dpi=200,
                    first_page=page_number,
                    last_page=page_number,
                    fmt="png",
                    poppler_path=poppler_path,
                    thread_count=1,
                )
                if not images:
                    raise LessonContentError(f"Không render được trang PDF {page_number}.")
                image = images[0]
                if image.width * image.height > MAX_IMAGE_PIXELS:
                    raise LessonContentError(f"Trang PDF {page_number} vượt giới hạn pixel OCR.")
                page_texts[page_number - 1] = normalize_text(
                    _ocr_image(_prepare_ocr_image(image))
                )
            except LessonContentError:
                raise
            except Exception as exc:
                raise LessonContentError(f"Không OCR được trang PDF {page_number}: {exc}") from exc

    return validate_text_length("\n\n".join(page_texts))


def extract_lesson_text(source: IncomingLesson) -> str:
    data = download_lesson(source)
    # Tin vào nội dung thật của tệp trước metadata do Zalo cung cấp. Cùng một
    # attachment có thể bị gắn nhãn FILE/IMAGE hoặc không có phần mở rộng.
    if data.startswith(b"%PDF-"):
        return _extract_pdf(data)
    if _is_docx(data):
        return _extract_docx(data)
    if _is_image(data):
        return _extract_image(data)
    return _extract_text_file(data)

"""Create the detailed DOCX report for the facial-expression benchmark."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "evaluation" / "results"
CSV_PATH = RESULTS / "facial_expression_benchmark.csv"
CHART_PATH = RESULTS / "facial_expression_benchmark.png"
OUTPUT_PATH = RESULTS / "bao_cao_benchmark_nhan_dien_bieu_cam.docx"

NAVY = "17365D"
BLUE = "2474E5"
PALE_BLUE = "EAF2FD"
PALE_GRAY = "F4F6F8"
WHITE = "FFFFFF"
BLACK = "000000"
GRAY = "5F6B7A"
BORDER = "D9D9D9"
GREEN = "1B7F63"
ORANGE = "C66B18"


def set_run_font(run, name="Arial", size=None, bold=None, color=BLACK, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, value, end):
        run._r.append(element)
    set_run_font(run, size=9, color=GRAY)


def suppress_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "between"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, *, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=10.5, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, size=10.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5)
    return paragraph


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def style_table(table, widths, *, header_fill=NAVY):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, WHITE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.02
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=9.2,
                        bold=(row_index == 0),
                        color=WHITE if row_index == 0 else BLACK,
                    )


def add_table(doc, headers, rows, widths, *, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    style_table(table, widths)
    if alignments:
        for row_index, row in enumerate(table.rows):
            for col_index, alignment in enumerate(alignments):
                row.cells[col_index].paragraphs[0].alignment = alignment
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    return table


def read_benchmark():
    with CSV_PATH.open("r", encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    return (
        [row for row in rows if row["section"] == "published"],
        [row for row in rows if row["section"] == "internal"],
        [row for row in rows if row["section"] == "distribution"],
    )


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)

    for level, size in ((1, 16), (2, 12.5), (3, 11)):
        style = styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(paragraph)

    doc.core_properties.title = "Báo cáo benchmark nhận diện biểu cảm khuôn mặt"
    doc.core_properties.subject = "Đánh giá mô hình và pipeline nhận diện biểu cảm của Zalo Health Bot"
    doc.core_properties.author = "Dự án Zalo Health Bot"
    doc.core_properties.keywords = "FER2013, facial expression recognition, benchmark, ONNX"


def add_title_page(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    suppress_paragraph_borders(title)
    run = title.add_run("Báo cáo benchmark nhận diện biểu cảm khuôn mặt")
    set_run_font(run, size=25, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Đánh giá mô hình, pipeline nội bộ và kế hoạch đo độ chính xác")
    set_run_font(run, size=13, color=GRAY)

    meta = doc.add_table(rows=4, cols=2)
    meta_data = [
        ("Hệ thống", "Zalo Health Bot"),
        ("Phạm vi", "Nhận diện bảy nhóm biểu cảm khuôn mặt và cơ chế từ chối unclear"),
        ("Ngày báo cáo", "14/10/2026"),
        ("Phiên bản", "1.0"),
    ]
    for row, (key, value) in zip(meta.rows, meta_data):
        row.cells[0].text = key
        row.cells[1].text = value
    style_table(meta, [1.35, 5.35], header_fill=NAVY)
    for row in meta.rows:
        set_cell_shading(row.cells[0], PALE_BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, size=9.5, bold=True, color=BLACK)
    doc.add_paragraph()

    add_heading(doc, "Tóm tắt điều hành", 1)
    add_body(
        doc,
        "Trọng số web-app đang được chuyển sang ONNX trong project có độ chính xác công bố 69,8% trên webcam test set. "
        "Kiểm tra nội bộ xác nhận pipeline trả kết quả có cấu trúc và định tuyến số khuôn mặt đúng trên 15/15 ảnh, "
        "nhưng chưa thể kết luận độ chính xác biểu cảm nội bộ vì tập ảnh này không có nhãn ground truth độc lập.",
    )
    add_body(doc, "Kết luận chính. Mô hình đủ cơ sở để tiếp tục thử nghiệm chức năng, nhưng chưa đủ bằng chứng để công bố accuracy end-to-end của chatbot trên người dùng thực tế.", bold_lead="Kết luận chính.")
    add_body(doc, "Khuyến nghị. Xây dựng tập kiểm thử có nhãn, tách theo người, báo cáo macro F1 và selective accuracy bên cạnh accuracy tổng thể trước khi dùng kết quả cho quyết định sản phẩm.", bold_lead="Khuyến nghị.")

    add_heading(doc, "Ba điểm cần ghi nhớ", 2)
    add_bullet(doc, "69,8% là số liệu công bố của mô hình web-app, không phải phép đo mới trên 15 ảnh nội bộ.")
    add_bullet(doc, "15/15 phản ánh tính ổn định và logic định tuyến của pipeline, không phải tỷ lệ nhận diện biểu cảm đúng.")
    add_bullet(doc, "Nhãn biểu cảm chỉ mô tả nhóm nét mặt; không chứng minh cảm xúc thật, tình trạng tâm lý hay chẩn đoán y khoa.")


def add_chart_page(doc):
    doc.add_page_break()
    add_heading(doc, "Tổng quan benchmark", 1)
    add_body(
        doc,
        "Biểu đồ dưới đây đặt số liệu công bố cạnh kiểm tra kỹ thuật nội bộ và phần chưa được đo. "
        "Cách trình bày này ngăn việc nhầm tỷ lệ trả kết quả hợp lệ với accuracy biểu cảm.",
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    picture = run.add_picture(str(CHART_PATH), width=Inches(6.75))
    doc_pr = picture._inline.docPr
    doc_pr.set("descr", "Biểu đồ benchmark nhận diện biểu cảm khuôn mặt với số liệu công bố, kiểm tra nội bộ và cấu hình")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run("Hình 1  Tổng quan benchmark nhận diện biểu cảm khuôn mặt")
    set_run_font(run, size=9.5, italic=True, color=GRAY)
    add_body(
        doc,
        "Cách đọc. Thanh 69,8% là accuracy công bố của trọng số web-app. Các chỉ số 15/15 và 4/4 chỉ đo khả năng pipeline tạo observation hợp lệ, phân biệt một mặt với nhiều mặt và chủ động trả unclear khi cần.",
        bold_lead="Cách đọc.",
    )


def add_system_page(doc):
    doc.add_page_break()
    add_heading(doc, "Mô hình và pipeline đang đánh giá", 1)
    add_body(
        doc,
        "Tính năng kết hợp phát hiện khuôn mặt OpenCV với mạng CNN năm lớp được chuyển từ trọng số TensorFlow.js của nhóm tác giả sang ONNX. "
        "Ảnh chỉ được phân loại khi bộ phát hiện tìm thấy đúng một khuôn mặt. Mọi trường hợp không có mặt, nhiều mặt hoặc xác suất không đủ rõ đều được đưa về unclear.",
    )
    add_heading(doc, "Luồng xử lý", 2)
    steps = [
        "Tiếp nhận ảnh và chuẩn hóa hướng ảnh theo metadata EXIF.",
        "Chuyển ảnh sang RGB để phát hiện khuôn mặt bằng Haar cascade của OpenCV.",
        "Yêu cầu đúng một khuôn mặt; cắt vùng mặt, chuyển ảnh xám và resize về 48 x 48 pixel.",
        "Chạy ONNX hai lần với ảnh gốc và ảnh lật ngang, sau đó lấy trung bình xác suất.",
        "Chỉ nhận top-1 khi confidence >= 0,50 và chênh lệch top-1 với top-2 >= 0,12; ngược lại trả unclear.",
        "Kết hợp nhãn nét mặt với dấu hiệu mắt và lời người dùng tự kể để tạo phản hồi có điều kiện.",
    ]
    for index, step in enumerate(steps, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.16)
        paragraph.paragraph_format.first_line_indent = Inches(-0.16)
        paragraph.paragraph_format.space_after = Pt(5)
        number = paragraph.add_run(f"{index}. ")
        set_run_font(number, size=10.5, bold=True, color=BLUE)
        text = paragraph.add_run(step)
        set_run_font(text, size=10.5)

    add_heading(doc, "Thông số vận hành", 2)
    add_table(
        doc,
        ["Thành phần", "Thông số", "Ý nghĩa"],
        [
            ("Đầu vào CNN", "48 x 48, grayscale", "Phù hợp kiến trúc FER2013 của trọng số gốc"),
            ("Số lớp", "7", "angry, disgust, fear, happy, sad, surprise, neutral"),
            ("Test-time augmentation", "2 lượt", "Ảnh gốc và lật ngang; trung bình xác suất"),
            ("Confidence tối thiểu", "0,50", "Loại dự đoán có xác suất top-1 thấp"),
            ("Margin tối thiểu", "0,12", "Loại dự đoán top-1 và top-2 quá gần nhau"),
            ("Quy tắc số mặt", "Đúng 1 mặt", "Không có hoặc nhiều mặt thì không gán nhãn"),
        ],
        [1.55, 1.55, 3.6],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )


def add_method_page(doc):
    doc.add_page_break()
    add_heading(doc, "Thiết kế benchmark và định nghĩa chỉ số", 1)
    add_heading(doc, "Hai nguồn bằng chứng", 2)
    add_body(
        doc,
        "Số liệu công bố đến từ báo cáo của Khanzada, Bai và Celepcikay năm 2020. Báo cáo nêu accuracy 69,8% và tốc độ nhận diện 40 ms cho mô hình web-app năm lớp trên webcam test set. Các mốc 66,3%, 71,2%, 75,2% và 75,8% dùng để đặt kết quả trong bối cảnh nghiên cứu.",
    )
    add_body(
        doc,
        "Kiểm tra nội bộ ngày 01/10/2026 dùng 15 ảnh thực tế qua cùng pipeline quan sát khuôn mặt. Tập gồm 11 ảnh một mặt và 4 ảnh nhiều mặt. Lần chạy cuối tạo observation hợp lệ cho 15/15 ảnh và đánh dấu unclear cho 4/4 ảnh nhiều mặt.",
    )

    add_heading(doc, "Định nghĩa chỉ số", 2)
    add_table(
        doc,
        ["Chỉ số", "Cách tính", "Cách diễn giải"],
        [
            ("Accuracy", "Dự đoán đúng / tổng mẫu có nhãn chuẩn", "Đo tỷ lệ đúng tổng thể; nhạy với mất cân bằng lớp"),
            ("Macro F1", "Trung bình F1 của bảy lớp", "Mỗi lớp có trọng số ngang nhau; quan trọng với lớp disgust ít mẫu"),
            ("Coverage", "Mẫu được gán nhãn / tổng mẫu", "Đo tỷ lệ hệ thống không trả unclear"),
            ("Selective accuracy", "Dự đoán đúng / số mẫu được gán nhãn", "Đo độ đúng sau khi hệ thống từ chối ca không chắc chắn"),
            ("Face-count routing", "Định tuyến đúng one, none hoặc unclear / tổng mẫu", "Đánh giá cổng an toàn trước phân loại biểu cảm"),
            ("Latency P95", "Phân vị 95 thời gian xử lý", "Phản ánh trải nghiệm trong tình huống chậm"),
        ],
        [1.35, 2.35, 3.0],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(
        doc,
        "Nguyên tắc báo cáo. Không dùng output validity, coverage hoặc abstention rate thay cho accuracy. Accuracy chỉ hợp lệ khi mỗi ảnh có nhãn ground truth độc lập và quy tắc tính được cố định trước khi chạy.",
        bold_lead="Nguyên tắc báo cáo.",
    )


def add_results_page(doc, published, internal, distribution):
    doc.add_page_break()
    add_heading(doc, "Kết quả chi tiết", 1)
    add_heading(doc, "Số liệu công bố và tham chiếu", 2)
    pub_rows = []
    role_names = {"published": "Công bố", "reference": "Tham chiếu"}
    dataset_names = {
        "FER2013 + auxiliary": "FER2013 + dữ liệu bổ trợ",
    }
    for row in published:
        pub_rows.append((
            row["metric"],
            dataset_names.get(row["dataset"], row["dataset"]),
            f'{float(row["value"]):.1f}%'.replace(".", ","),
            role_names.get(row["status"], row["status"]),
        ))
    add_table(
        doc,
        ["Mô hình hoặc mốc", "Tập đánh giá", "Accuracy", "Vai trò"],
        pub_rows,
        [2.2, 2.25, 0.9, 1.35],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_body(
        doc,
        "Không nên xếp hạng trực tiếp 69,8% với mọi thanh còn lại vì webcam test set và FER2013 private test không phải cùng một phân phối. Giá trị 69,8% phù hợp nhất để mô tả trọng số web-app đang dùng; các giá trị khác chỉ là mốc tham chiếu.",
    )

    add_heading(doc, "Kiểm tra pipeline nội bộ", 2)
    int_rows = []
    internal_dataset_names = {
        "15 local photos": "15 ảnh nội bộ",
        "4 multi-face photos": "4 ảnh nhiều mặt",
    }
    for row in internal:
        result = "N/A" if row["status"] == "not_measured" else f'{row["numerator"]}/{row["denominator"]} ({float(row["value"]):.0f}%)'
        int_rows.append((row["metric"], internal_dataset_names.get(row["dataset"], row["dataset"]), result, "Chưa đo" if row["status"] == "not_measured" else "Đã đo"))
    add_table(
        doc,
        ["Tiêu chí", "Tập kiểm tra", "Kết quả", "Trạng thái"],
        int_rows,
        [2.25, 2.05, 1.3, 1.1],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "Phân bố nhãn dự đoán", 2)
    dist_rows = []
    total = sum(int(float(row["value"])) for row in distribution)
    for row in distribution:
        count = int(float(row["value"]))
        dist_rows.append((row["metric"], count, f"{count / total:.1%}".replace(".", ",")))
    add_table(
        doc,
        ["Nhãn", "Số ảnh", "Tỷ trọng trong 11 ảnh một mặt"],
        dist_rows,
        [2.5, 1.2, 3.0],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )


def add_interpretation_page(doc):
    doc.add_page_break()
    add_heading(doc, "Diễn giải và giới hạn", 1)
    add_heading(doc, "Điều có thể kết luận", 2)
    add_bullet(doc, "Mô hình web-app có accuracy công bố 69,8% trên tập webcam test của tác giả và đáp ứng mục tiêu tốc độ công bố 40 ms trong môi trường web-app.")
    add_bullet(doc, "Pipeline nội bộ đã xử lý ổn định 15 ảnh kiểm tra, gồm cả trường hợp có người ở hậu cảnh.")
    add_bullet(doc, "Cơ chế confidence, margin và unclear làm giảm nguy cơ gán nhãn khi mô hình không chắc chắn.")

    add_heading(doc, "Điều chưa thể kết luận", 2)
    add_bullet(doc, "Chưa có accuracy, precision, recall hoặc F1 đo trực tiếp trên ảnh người dùng của Zalo Health Bot.")
    add_bullet(doc, "Không thể suy ra cảm xúc thật, mức độ căng thẳng, trầm cảm hoặc nguy cơ sức khỏe tâm thần từ nhãn nét mặt.")
    add_bullet(doc, "Không thể dùng 15/15 output hợp lệ để khẳng định mô hình nhận diện đúng 100%.")

    add_heading(doc, "Các nguồn sai số chính", 2)
    add_table(
        doc,
        ["Nguồn sai số", "Tác động", "Biện pháp kiểm soát"],
        [
            ("Ánh sáng, góc nghiêng, che mặt", "Giảm chất lượng phát hiện và crop mặt", "Thu thập test theo điều kiện; trả unclear khi không đủ rõ"),
            ("Mất cân bằng lớp", "Accuracy tổng che giấu lớp khó như disgust và fear", "Báo cáo macro F1 và recall theo lớp"),
            ("Dataset shift", "Webcam test khác ảnh điện thoại và người dùng Việt Nam", "Tạo tập kiểm thử đại diện cho kênh Zalo"),
            ("Nhãn cảm xúc chủ quan", "Hai người có thể gán nhãn khác nhau", "Hai annotator độc lập và một người phân xử"),
            ("Face detector Haar", "Có thể bỏ sót mặt nghiêng, nhỏ hoặc nhiều mặt", "Đánh giá face-count routing riêng khỏi classifier"),
            ("Ngưỡng từ chối", "Tăng độ đúng có thể làm giảm coverage", "Vẽ đường selective accuracy theo coverage"),
        ],
        [1.65, 2.25, 2.8],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )


def add_plan_page(doc):
    doc.add_page_break()
    add_heading(doc, "Kế hoạch benchmark nội bộ có ground truth", 1)
    add_heading(doc, "Thiết kế tập dữ liệu", 2)
    add_bullet(doc, "Thu thập tối thiểu 700 ảnh hợp lệ, mục tiêu ít nhất 100 ảnh cho mỗi lớp sau kiểm duyệt.")
    add_bullet(doc, "Có ít nhất 50 người tham gia và tách train, validation, test theo người để tránh cùng một khuôn mặt xuất hiện ở nhiều tập.")
    add_bullet(doc, "Bao phủ ánh sáng tốt và yếu, chính diện và nghiêng, kính, khẩu trang, nền đông người và camera có chất lượng khác nhau.")
    add_bullet(doc, "Mỗi ảnh được hai người gán nhãn độc lập; bất đồng do người thứ ba phân xử. Báo cáo mức đồng thuận Cohen kappa hoặc Fleiss kappa.")
    add_bullet(doc, "Tách riêng bộ one-face để đo classifier và bộ none hoặc multi-face để đo face-count routing.")

    add_heading(doc, "Bộ chỉ số bắt buộc", 2)
    add_table(
        doc,
        ["Nhóm", "Chỉ số", "Mục đích"],
        [
            ("Classifier", "Accuracy, macro F1, weighted F1", "Đo tổng thể và cân bằng giữa các lớp"),
            ("Theo lớp", "Precision, recall, F1, support", "Nhìn rõ lớp yếu và ảnh hưởng mất cân bằng"),
            ("Từ chối", "Coverage, abstention rate, selective accuracy", "Đánh giá hiệu quả của confidence và margin"),
            ("Định tuyến", "One/none/multi-face confusion matrix", "Đo cổng an toàn trước classifier"),
            ("Hiệu năng", "Latency P50/P95, lỗi inference, memory", "Đánh giá khả năng vận hành"),
            ("Công bằng", "Chỉ số theo điều kiện chụp và nhóm người", "Phát hiện chênh lệch hiệu năng"),
        ],
        [1.25, 2.45, 3.0],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "Ngưỡng nghiệm thu đề xuất", 2)
    add_body(doc, "Các ngưỡng sau là mục tiêu đề xuất cho vòng đánh giá đầu tiên, không phải kết quả hiện tại:")
    add_bullet(doc, "Face-count routing accuracy >= 95% trên tập none, one và multi-face cân bằng.")
    add_bullet(doc, "Structured output validity >= 99%.")
    add_bullet(doc, "Macro F1 >= 0,65 và không lớp nào có recall dưới 0,45.")
    add_bullet(doc, "Selective accuracy >= 75% tại coverage >= 70%, sau khi cố định ngưỡng trên validation set.")
    add_bullet(doc, "Latency P95 được đo và chấp nhận theo môi trường triển khai thực tế; không dùng tốc độ 40 ms công bố thay cho phép đo hệ thống hiện tại.")


def add_reproducibility_page(doc):
    doc.add_page_break()
    add_heading(doc, "Tái tạo kết quả và tài liệu tham chiếu", 1)
    add_heading(doc, "Tệp và lệnh trong project", 2)
    add_body(doc, "Biểu đồ và bảng dữ liệu của báo cáo được lưu trong thư mục evaluation/results. Có thể tái tạo biểu đồ bằng lệnh:")
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.space_after = Pt(9)
    run = paragraph.add_run(r".\.venv\Scripts\python.exe evaluation\facial_benchmark_chart.py")
    set_run_font(run, name="Consolas", size=9.5, color=NAVY)

    add_body(doc, "Để chạy pipeline quan sát khuôn mặt trên một thư mục ảnh local:")
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.space_after = Pt(9)
    run = paragraph.add_run(r".\.venv\Scripts\python.exe -m evaluation.facial_image_eval <thu_muc_anh>")
    set_run_font(run, name="Consolas", size=9.5, color=NAVY)

    add_heading(doc, "Tài liệu tham chiếu", 2)
    references = [
        (
            "Khanzada A, Bai C, Celepcikay FT. Facial Expression Recognition with Deep Learning. Stanford CS230, 2020.",
            "https://cs230.stanford.edu/projects_winter_2020/reports/32610274.pdf",
            "Báo cáo Stanford CS230",
        ),
        (
            "Goodfellow IJ và cộng sự. Challenges in Representation Learning. FER2013 benchmark, 2013.",
            "https://doi.org/10.1007/978-3-642-42051-1_16",
            "DOI bài báo FER2013",
        ),
        (
            "Tài liệu kỹ thuật nội bộ FACIAL_SUPPORT.md, lần kiểm tra ảnh ngày 01/10/2026.",
            "",
            "",
        ),
    ]
    for index, (citation, url, link_text) in enumerate(references, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.23)
        paragraph.paragraph_format.first_line_indent = Inches(-0.23)
        paragraph.paragraph_format.space_after = Pt(7)
        run = paragraph.add_run(f"{index}. {citation} ")
        set_run_font(run, size=10)
        if url:
            add_hyperlink(paragraph, link_text, url)

    add_heading(doc, "Kết luận", 2)
    add_body(
        doc,
        "Benchmark hiện tại cung cấp bằng chứng tốt về khả năng vận hành của pipeline và một mốc accuracy công bố phù hợp với trọng số gốc. "
        "Bước tiếp theo có giá trị cao nhất là xây dựng tập test có ground truth đại diện cho người dùng và điều kiện chụp thực tế của Zalo Health Bot. "
        "Chỉ sau phép đo đó mới nên công bố accuracy end-to-end của tính năng.",
    )


def build():
    published, internal, distribution = read_benchmark()
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_chart_page(doc)
    add_system_page(doc)
    add_method_page(doc)
    add_results_page(doc, published, internal, distribution)
    add_interpretation_page(doc)
    add_plan_page(doc)
    add_reproducibility_page(doc)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    build()
